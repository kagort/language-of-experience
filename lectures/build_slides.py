"""Сборка слайдов Лекции III: python3 build_slides.py → lecture3-slides.pptx + outline."""
import io
import re
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
IMG = ROOT / "assets" / "images" / "philosophers"

C = {k: RGBColor.from_string(v) for k, v in {
    "bg": "F2E4C1", "card": "EDE0C0", "line": "D6C48A", "ink": "1A0D05",
    "dark": "3B1A08", "mid": "6B3D1E", "light": "9C6240", "gold": "8B6914",
    "goldl": "C9A227", "pale": "E8D4A0", "white": "FFFFFF",
    "western": "2E5F8F", "patristic": "9B2D24", "hindu": "C76A12",
    "buddhist": "3D7A55", "shared": "6E5A45",
}.items()}
TRAD_RU = {"western": "Западная философия", "patristic": "Святоотеческая аскетика",
           "hindu": "Индуизм", "buddhist": "Буддизм"}
HEAD, BODY = "Playfair Display", "EB Garamond"

PORTRAIT = {
    "patanjali": "patanjali.jpg", "shankara": "shankara.jpg", "ramanuja": "ramanuja.jpg",
    "ramana_maharshi": "ramana_maharshi.jpg", "bhartrhari": "bhartrhari.jpg",
    "abhinavagupta": "abhinavagupta.webp", "nagarjuna": "nagarjuna.jpeg",
    "chandrakirti": "chandrakirti.jpg", "nagasena": "nagasena.jpg", "buddhaghosa": None,
    "vasubandhu": "vasubandhu.jpg", "asanga": "asanga.jpg", "dignaga": "dignaga.jpg",
    "dharmakirti": "dharmakirti.jpg", "linji_yixuan": "linji_yixuan.jpg",
    "tilopa": "tilopa.jpg", "dogen": "dogen.png", "nisargadatta": "nisargadatta_maharaj.jpg",
}
NAMES = {
    "patanjali": "Патанджали", "shankara": "Шанкара", "ramanuja": "Рамануджа",
    "ramana_maharshi": "Рамана Махарши", "bhartrhari": "Бхартрихари",
    "abhinavagupta": "Абхинавагупта", "nagarjuna": "Нагарджуна", "chandrakirti": "Чандракирти",
    "nagasena": "Нагасена", "buddhaghosa": "Буддхагхоша", "vasubandhu": "Васубандху",
    "asanga": "Асанга", "dignaga": "Дигнага", "dharmakirti": "Дхармакирти",
    "linji_yixuan": "Линьцзи", "tilopa": "Тилопа", "dogen": "Догэн",
    "nisargadatta": "Нисаргадатта",
}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
OUTLINE = []


# ── primitives ──────────────────────────────────────────────
def rect(s, x, y, w, h, fill=None, line=None, lw=1, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, x, y, w, h)
    if fill is None:
        r.fill.background()
    else:
        r.fill.solid(); r.fill.fore_color.rgb = C[fill] if isinstance(fill, str) else fill
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = C[line]; r.line.width = Pt(lw)
    r.shadow.inherit = False
    return r


def text(s, x, y, w, h, runs, size=20, color="ink", font=BODY, bold=False, italic=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.1):
    """runs: str | list of paragraphs; paragraph = str | list of (text, {opts})."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        parts = para if isinstance(para, list) else [(para, {})]
        for t, o in parts:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = C[o.get("color", color)]
    return tb


def base(title=None, section=None, kicker=None):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, H, fill="bg")
    rect(s, Inches(0.35), Inches(0.35), W - Inches(0.7), H - Inches(0.7), line="line", lw=0.75)
    if section:
        text(s, Inches(0.7), Inches(6.72), Inches(8), Inches(0.35), section, size=11,
             color="light", font=BODY, italic=True)
    if title:
        if kicker:
            text(s, Inches(0.7), Inches(0.55), Inches(11.5), Inches(0.35), kicker.upper(), size=12,
                 color="gold", font=BODY, bold=True)
        text(s, Inches(0.7), Inches(0.8 if kicker else 0.6), Inches(11.9), Inches(0.9), title,
             size=34, color="dark", font=HEAD, bold=True)
        rect(s, Inches(0.72), Inches(1.72 if kicker else 1.55), Inches(1.2), Emu(22000), fill="goldl")
    return s


def notes(s, n, title):
    OUTLINE.append((n, title))
    s.notes_slide.notes_text_frame.text = NOTES.get(n, "")


def portrait(s, pid, x, y, w, label=True):
    h = int(w * 1.2)
    rect(s, x - Emu(30000), y - Emu(30000), w + Emu(60000), h + Emu(60000), fill="card", line="gold", lw=1)
    f = PORTRAIT.get(pid)
    if f:
        im = Image.open(IMG / f).convert("RGB")
        tw, th = im.size
        tr = 5 / 6
        if tw / th > tr:
            nw = int(th * tr); im = im.crop(((tw - nw) // 2, 0, (tw - nw) // 2 + nw, th))
        else:
            nh = int(tw / tr); top = max(0, int((th - nh) * 0.25)); im = im.crop((0, top, tw, top + nh))
        im = im.resize((500, 600))
        buf = io.BytesIO(); im.save(buf, "JPEG", quality=85); buf.seek(0)
        s.shapes.add_picture(buf, x, y, w, h)
    else:
        rect(s, x, y, w, h, fill="line")
        text(s, x, y, w, h, NAMES[pid][0], size=int(w / 12700 / 1.6), color="mid", font=HEAD,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if label:
        text(s, x - Inches(0.3), y + h + Inches(0.08), w + Inches(0.6), Inches(0.35), NAMES[pid],
             size=13, color="mid", align=PP_ALIGN.CENTER)


def card(s, x, y, w, h, head, body, color="gold", head_size=19, body_size=16, fill="card"):
    rect(s, x, y, w, h, fill=fill, line="line", lw=0.75)
    rect(s, x, y, Inches(0.09), h, fill=color)
    text(s, x + Inches(0.25), y + Inches(0.14), w - Inches(0.4), Inches(0.5), head, size=head_size,
         color=color if color in C and color not in ("gold",) else "dark", font=HEAD, bold=True)
    text(s, x + Inches(0.25), y + Inches(0.14) + Pt(head_size * 1.9), w - Inches(0.4),
         h - Inches(0.3) - Pt(head_size * 1.9), body, size=body_size, color="ink")


def quote(s, q, src, y=Inches(2.1), size=30, w=Inches(10.8)):
    x = (W - w) // 2
    text(s, x - Inches(0.6), y - Inches(0.55), Inches(1), Inches(1.2), "«", size=90, color="goldl", font=HEAD)
    text(s, x, y, w, Inches(3), q, size=size, color="dark", font=HEAD, italic=True, spacing=1.15)
    text(s, x, y + Inches(3.25), w, Inches(0.5), "— " + src, size=16, color="light", align=PP_ALIGN.RIGHT)


def table(s, x, y, w, rows, colw, header_colors=None, size=14, head_size=15, rowh=Inches(0.6)):
    n_rows, n_cols = len(rows), len(rows[0])
    shp = s.shapes.add_table(n_rows, n_cols, x, y, w, rowh * n_rows)
    tbl = shp.table
    tbl.first_row = True
    for i, cw in enumerate(colw):
        tbl.columns[i].width = cw
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.fill.solid()
            if r == 0:
                hc = header_colors[c] if header_colors else "dark"
                cell.fill.fore_color.rgb = C[hc] if hc else C["bg"]
            else:
                cell.fill.fore_color.rgb = C["card"] if r % 2 else C["bg"]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run(); run.text = val
            run.font.name = HEAD if (r == 0 or c == 0) else BODY
            run.font.size = Pt(head_size if r == 0 else size)
            run.font.bold = r == 0 or c == 0
            run.font.color.rgb = C["white"] if (r == 0 and (not header_colors or header_colors[c])) else C["dark" if c == 0 else "ink"]
    return tbl


def divider(n, roman, title, subtitle, people, section):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, H, fill="dark")
    rect(s, Inches(0.35), Inches(0.35), W - Inches(0.7), H - Inches(0.7), line="gold", lw=0.75)
    text(s, Inches(0.9), Inches(1.0), Inches(7), Inches(0.5), f"РАЗДЕЛ {roman}", size=14, color="goldl", bold=True)
    tw = Inches(11.5) if not people else (Inches(5.6) if len(people) > 3 else Inches(6.6))
    text(s, Inches(0.9), Inches(1.45), tw, Inches(2.2), title, size=40, color="white", font=HEAD, bold=True)
    text(s, Inches(0.9), Inches(3.75), tw, Inches(1.4), subtitle, size=20, color="pale", italic=True)
    text(s, Inches(0.9), Inches(6.0), Inches(1), Inches(0.6), "❦", size=26, color="goldl", font=HEAD)
    if people:
        pw = Inches(1.45) if len(people) <= 3 else Inches(1.15)
        gap = Inches(0.3)
        total = len(people) * pw + (len(people) - 1) * gap
        x0 = Inches(12.4) - total
        for i, pid in enumerate(people):
            x = x0 + i * (pw + gap)
            portrait(s, pid, x, Inches(2.5), pw, label=False)
            text(s, x - Inches(0.3), Inches(2.5) + int(pw * 1.2) + Inches(0.1), pw + Inches(0.6),
                 Inches(0.4), NAMES[pid], size=13, color="pale", align=PP_ALIGN.CENTER)
    notes(s, n, f"Раздел {roman}. {title}")
    return s


# ── notes from script ───────────────────────────────────────
script = (HERE / "lecture3-script.md").read_text(encoding="utf-8")
NOTES = {}
parts = re.split(r"\[Слайд (\d+)\]", script)
for i in range(1, len(parts), 2):
    chunk = re.split(r"\n## ", parts[i + 1])[0]
    chunk = re.sub(r"\*\*(.+?)\*\*", r"\1", chunk)
    chunk = re.sub(r"(?m)^---\s*$", "", chunk)
    NOTES[int(parts[i])] = re.sub(r"\n{3,}", "\n\n", chunk).strip()

S1 = "I · Один вопрос — четыре ответа"
S2 = "II · Веданта и йога"
S3 = "III · Бхартрихари"
S4 = "IV · Мадхьямака"
S5 = "V · Йогачара"
S6 = "VI · Сравнение"
S7 = "VII · Молчание и практика"
S8 = "VIII · Синтез цикла"

# ── 1. Title ────────────────────────────────────────────────
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill="bg")
rect(s, Inches(0.35), Inches(0.35), W - Inches(0.7), H - Inches(0.7), line="gold", lw=1)
rect(s, Inches(0.48), Inches(0.48), W - Inches(0.96), H - Inches(0.96), line="line", lw=0.5)
text(s, 0, Inches(1.0), W, Inches(0.7), "❦", size=40, color="gold", font=HEAD, align=PP_ALIGN.CENTER)
text(s, 0, Inches(1.75), W, Inches(0.4), "СЕРИЯ ЛЕКЦИЙ «СОЗНАНИЕ · ТЕЛО · ЯЗЫК» · ЛЕКЦИЯ III", size=14,
     color="gold", bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(2.3), W - Inches(2), Inches(1.2), "Сознание, реальность и язык", size=54,
     color="dark", font=HEAD, bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(1.5), Inches(3.55), W - Inches(3), Inches(1.0),
     "Академическая философия, отцы Церкви и учения Востока", size=26, color="mid", font=HEAD,
     italic=True, align=PP_ALIGN.CENTER)
rect(s, (W - Inches(1.6)) // 2, Inches(4.75), Inches(1.6), Emu(22000), fill="goldl")
text(s, 0, Inches(5.0), W, Inches(0.5), "Заключительная лекция цикла", size=18, color="light",
     align=PP_ALIGN.CENTER)
notes(s, 1, "Титульный слайд")

# ── 2. Cycle map ────────────────────────────────────────────
s = base("Три вечера — один вопрос", S1, "Карта цикла")
text(s, Inches(0.7), Inches(1.85), Inches(11.9), Inches(0.6),
     "На что указывает слово, когда мы говорим о чувственном и внутреннем опыте?", size=20,
     color="mid", italic=True)
cols = [("I", "Западная мысль", "Декарт → Кант → Гуссерль → Витгенштейн → Барретт, Фристон",
         "Язык — конструктор опыта", "western"),
        ("II", "Отцы-аскеты", "Евагрий, Лествичник, Исаак Сирин, Симеон, Феофан",
         "Расхождения важнее параллелей", "patristic"),
        ("III", "Индия и буддизм", "Патанджали, Шанкара, Бхартрихари, Нагарджуна, Васубандху, Догэн",
         "Что видно с четырёх сторон?", "buddhist")]
for i, (r, h, b, t, c) in enumerate(cols):
    x = Inches(0.7) + i * Inches(4.05)
    rect(s, x, Inches(2.75), Inches(3.75), Inches(3.6), fill="card", line="line")
    rect(s, x, Inches(2.75), Inches(3.75), Inches(0.09), fill=c)
    text(s, x + Inches(0.25), Inches(2.95), Inches(1), Inches(0.8), r, size=40, color=c, font=HEAD, bold=True)
    text(s, x + Inches(0.25), Inches(3.75), Inches(3.3), Inches(0.5), h, size=22, color="dark", font=HEAD, bold=True)
    text(s, x + Inches(0.25), Inches(4.3), Inches(3.3), Inches(1.0), b, size=15, color="light")
    text(s, x + Inches(0.25), Inches(5.35), Inches(3.3), Inches(0.9), t, size=18, color="ink", italic=True)
notes(s, 2, "Карта цикла: три лекции")

# ── 3. Where we stopped ─────────────────────────────────────
s = base("Где мы остановились", S1, "Итоги лекций I и II")
for i, (k, q, c) in enumerate([
        ("Лекция I", "Язык не зеркало опыта, а один из его конструкторов.", "western"),
        ("Лекция II", "Параллели реальны, но расхождения важнее параллелей.", "patristic")]):
    y = Inches(2.2) + i * Inches(2.05)
    rect(s, Inches(0.9), y, Inches(0.09), Inches(1.6), fill=c)
    text(s, Inches(1.2), y, Inches(10), Inches(0.4), k.upper(), size=14, color=c, bold=True)
    text(s, Inches(1.2), y + Inches(0.4), Inches(11), Inches(1.2), q, size=32, color="dark", font=HEAD, italic=True)
text(s, Inches(1.2), Inches(6.2), Inches(11), Inches(0.5),
     "Сегодня: ещё два голоса — индийская мысль и буддизм", size=18, color="light")
notes(s, 3, "Итоги лекций I и II")

# ── 4. Four languages ───────────────────────────────────────
s = base("Четыре языка для одного вопроса", S1)
langs = [("Западная философия", "язык понятий", "Что такое сознание?", "western"),
         ("Святоотеческая аскетика", "язык практики", "Как его исцелить?", "patristic"),
         ("Индуизм: Веданта и йога", "язык бытия", "Как из сознания возникает мир?", "hindu"),
         ("Буддизм", "язык анализа", "Кто спрашивает?", "buddhist")]
for i, (h, sub, q, c) in enumerate(langs):
    x = Inches(0.7) + (i % 2) * Inches(6.05)
    y = Inches(2.0) + (i // 2) * Inches(2.3)
    rect(s, x, y, Inches(5.8), Inches(2.0), fill="card", line="line")
    rect(s, x, y, Inches(0.1), Inches(2.0), fill=c)
    text(s, x + Inches(0.35), y + Inches(0.2), Inches(5.2), Inches(0.5), h, size=22, color=c, font=HEAD, bold=True)
    text(s, x + Inches(0.35), y + Inches(0.72), Inches(5.2), Inches(0.4), sub, size=16, color="light", italic=True)
    text(s, x + Inches(0.35), y + Inches(1.15), Inches(5.3), Inches(0.7), f"«{q}»", size=21, color="dark", font=HEAD)
notes(s, 4, "Четыре языка: понятия, практика, бытие, анализ")

# ── 5. Method ───────────────────────────────────────────────
s = base("Метод: параллельное чтение", S1)
card(s, Inches(0.7), Inches(1.95), Inches(3.8), Inches(1.7), "Ловушка 1",
     "«Все традиции говорят одно и то же»", "patristic", body_size=17)
card(s, Inches(0.7), Inches(3.85), Inches(3.8), Inches(1.7), "Ловушка 2",
     "«Традиции несоизмеримы, сравнивать нельзя»", "patristic", body_size=17)
text(s, Inches(0.7), Inches(5.75), Inches(3.8), Inches(0.9),
     "Слепые и слон («Удана»): каждый принимает свою часть за целое", size=14, color="light", italic=True)
text(s, Inches(5.0), Inches(1.95), Inches(7.5), Inches(0.5), "Четыре вопроса на весь вечер", size=22,
     color="gold", font=HEAD, bold=True)
qs = ["Сознание — основа реальности или её продукт?",
      "Восприятие раскрывает реальность или конструирует?",
      "Какова роль тела?",
      "Что язык делает с опытом?"]
for i, q in enumerate(qs):
    y = Inches(2.65) + i * Inches(0.95)
    text(s, Inches(5.0), y, Inches(0.7), Inches(0.7), str(i + 1), size=30, color="goldl", font=HEAD, bold=True)
    text(s, Inches(5.7), y + Inches(0.08), Inches(7), Inches(0.8), q, size=22, color="dark")
notes(s, 5, "Метод: две ловушки и четыре вопроса")

# ── II. Vedanta ─────────────────────────────────────────────
divider(6, "II", "Веданта и йога", "Сознание как основа: дано сознание, объяснения требует мир",
        ["patanjali", "shankara", "ramana_maharshi"], S2)

s = base(None, S2)
text(s, 0, Inches(1.2), W, Inches(0.5), "ЙОГА-СУТРЫ 1.2", size=14, color="gold", bold=True, align=PP_ALIGN.CENTER)
text(s, 0, Inches(1.8), W, Inches(0.9), "yogaś citta-vṛtti-nirodhaḥ", size=40, color="hindu", font=HEAD,
     italic=True, align=PP_ALIGN.CENTER)
text(s, 0, Inches(2.75), W, Inches(0.8), "Йога есть прекращение волнений читты", size=30, color="dark",
     font=HEAD, align=PP_ALIGN.CENTER)
for i, (t, d) in enumerate([("читта", "всё поле психической жизни"), ("вритти", "«водовороты»: пять видов движений"),
                            ("ниродха", "успокоение, прекращение")]):
    x = Inches(1.0) + i * Inches(3.85)
    rect(s, x, Inches(4.2), Inches(3.5), Inches(1.7), fill="card", line="line")
    text(s, x, Inches(4.4), Inches(3.5), Inches(0.6), t, size=26, color="hindu", font=HEAD, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + Inches(0.2), Inches(5.05), Inches(3.1), Inches(0.8), d, size=17, color="ink", align=PP_ALIGN.CENTER)
notes(s, 7, "Читта-вритти-ниродха (ЙС 1.2)")

s = base("Где проходит граница?", S2, "Санкхья и Декарт")
for i, (h, left, right, c) in enumerate([
        ("Декарт", "мышление · чувство · воля · память", "протяжённое тело", "western"),
        ("Санкхья", "Пуруша — чистый свидетель", "Пракрити: тело и читта — мысли, память, «я»", "hindu")]):
    y = Inches(2.1) + i * Inches(1.9)
    text(s, Inches(0.7), y + Inches(0.35), Inches(2.0), Inches(0.6), h, size=24, color=c, font=HEAD, bold=True)
    rect(s, Inches(2.8), y, Inches(4.4), Inches(1.4), fill="card", line=c, lw=1.25)
    text(s, Inches(2.95), y + Inches(0.05), Inches(4.1), Inches(0.35), "СОЗНАНИЕ", size=11, color="light", bold=True)
    text(s, Inches(2.95), y + Inches(0.4), Inches(4.1), Inches(1.0), left, size=19, color="dark")
    rect(s, Inches(7.4), y, Inches(5.2), Inches(1.4), fill="bg", line="line")
    text(s, Inches(7.55), y + Inches(0.05), Inches(4.9), Inches(0.35), "МАТЕРИЯ", size=11, color="light", bold=True)
    text(s, Inches(7.55), y + Inches(0.4), Inches(4.9), Inches(1.0), right, size=19, color="dark")
text(s, Inches(0.7), Inches(5.95), Inches(12), Inches(0.6),
     "Почему работа ума кому-то «светит»? — та же черта, что у Чалмерса между «лёгкими» и «трудной» проблемой",
     size=17, color="mid", italic=True)
notes(s, 8, "Санкхья vs Декарт; озеро; помыслы Евагрия")

s = base("Шанкара: Адвайта — «не-два»", S2)
card(s, Inches(0.7), Inches(1.95), Inches(5.8), Inches(1.55), "Брахман = сат · чит · ананда",
     "бытие · сознание · блаженство", "hindu", body_size=20)
card(s, Inches(0.7), Inches(3.7), Inches(5.8), Inches(1.55), "Атман = Брахман",
     "tat tvam asi — «то есть ты» (Чхандогья-упанишада)", "hindu", body_size=20)
rect(s, Inches(6.9), Inches(1.95), Inches(5.7), Inches(4.35), fill="card", line="line")
text(s, Inches(7.15), Inches(2.1), Inches(5.3), Inches(0.5), "Змея и верёвка", size=24, color="dark", font=HEAD, bold=True)
text(s, Inches(7.15), Inches(2.7), Inches(5.3), Inches(3.4),
     ["В сумерках — змея. Страх настоящий.", "С фонарём — верёвка.", "",
      [("Майя ", {"bold": True, "color": "hindu"}), ("— наложение (адхьяса): мир реален как опыт, но не как нечто независимое от сознания.", {})]],
     size=19, color="ink")
text(s, Inches(0.7), Inches(5.5), Inches(5.8), Inches(0.9), "«Трудная проблема» не решается — переворачивается",
     size=17, color="mid", italic=True)
notes(s, 9, "Шанкара: сат-чит-ананда, tat tvam asi, майя")

s = base("Четыре состояния и «четвёртое»", S2, "Мандукья-упанишада")
for i, (t, d) in enumerate([("джаграт", "бодрствование"), ("свапна", "сон со сновидениями"),
                            ("сушупти", "сон без сновидений")]):
    x = Inches(1.3) + i * Inches(3.7)
    rect(s, x, Inches(2.75), Inches(3.3), Inches(1.35), fill="card", line="line")
    text(s, x, Inches(2.85), Inches(3.3), Inches(0.55), t, size=24, color="hindu", font=HEAD, bold=True, align=PP_ALIGN.CENTER)
    text(s, x, Inches(3.4), Inches(3.3), Inches(0.5), d, size=16, color="ink", align=PP_ALIGN.CENTER)
rect(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(2.4), line="hindu", lw=2,
     shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(1.1), Inches(2.25), Inches(6), Inches(0.45), "ТУРЬЯ — свидетель всех трёх", size=14, color="hindu", bold=True)
text(s, Inches(0.9), Inches(4.95), Inches(11.5), Inches(1.4),
     ["Аргумент глубокого сна: «я спал и ничего не помню» — значит, кто-то это помнит.",
      [("Декарт: «мыслю — существую». Адвайта: «осознаётся — значит, есть осознающее».", {"italic": True, "color": "mid"})]],
     size=19, color="ink")
notes(s, 10, "Турья и аргумент глубокого сна")

s = base("Рамануджа и Рамана Махарши", S2, "Два пути после Шанкары")
portrait(s, "ramanuja", Inches(0.9), Inches(2.1), Inches(1.8))
text(s, Inches(3.0), Inches(2.1), Inches(3.4), Inches(3.6),
     [[("Вишишта-адвайта", {"bold": True, "color": "hindu", "font": HEAD, "size": 22})],
      "Мир и души реальны — «тело» Бога.", "Путь бхакти: познание как встреча.",
      [("Ближе к отцам, чем Шанкара", {"italic": True, "color": "mid"})]], size=18)
portrait(s, "ramana_maharshi", Inches(6.9), Inches(2.1), Inches(1.8))
text(s, Inches(9.0), Inches(2.1), Inches(3.6), Inches(3.6),
     [[("«Кто я?»", {"bold": True, "color": "hindu", "font": HEAD, "size": 26})],
      "Кому пришла эта мысль? — Мне.", "А кто я?",
      [("Внимание — от содержаний к источнику", {"italic": True, "color": "mid"})]], size=18)
text(s, Inches(0.9), Inches(5.85), Inches(11.6), Inches(0.6),
     "Витгенштейн: чем проверить «чистое я»? — Традиции: учитель, текст, плоды жизни", size=16, color="light", italic=True)
notes(s, 11, "Рамануджа, Рамана, возражение Витгенштейна")

s = base("Три «я»", S2, "Параллели и расхождения")
table(s, Inches(0.7), Inches(2.0), Inches(11.9), [
    ["", "Гуссерль", "Шанкара", "Отцы-аскеты"],
    ["Что остаётся", "трансцендентальное эго", "Атман", "нус, сердце"],
    ["Статус", "функция описания", "абсолют = Брахман", "сотворённая личность"],
    ["Личное?", "вопрос в скобках", "растворяется", "сохраняется"],
    ["Обращено к", "миру как явлению", "себе как Брахману", "Другому — Богу"],
], [Inches(2.3), Inches(3.2), Inches(3.2), Inches(3.2)],
    header_colors=[None, "western", "hindu", "patristic"], size=18, head_size=18, rowh=Inches(0.75))
notes(s, 12, "Таблица трёх «я»")

# ── III. Bhartrhari ─────────────────────────────────────────
divider(13, "III", "Бхартрихари", "Язык как первичная реальность", ["bhartrhari", "abhinavagupta"], S3)

s = base(None, S3)
quote(s, "Безначальный и бесконечный Брахман, чья сущность — Слово, развёртывается в виде вещей; из него — весь процесс мира",
      "«Вакьяпадия» I.1, пересказ", y=Inches(1.3), size=28)
text(s, Inches(1.25), Inches(5.3), Inches(10.8), Inches(0.9),
     [[("Нет в мире познания без слова. ", {"bold": True, "color": "hindu"}), ("Всякое знание как будто сплетено со словом. (I.123)", {})]],
     size=20, color="ink")
notes(s, 14, "Шабда-брахман: Вакьяпадия I.1 и I.123")

s = base("Спхота — смысл вспыхивает целиком", S3)
for i, syl in enumerate(["ко", "ро", "ва"]):
    x = Inches(1.0) + i * Inches(1.7)
    rect(s, x, Inches(2.3), Inches(1.4), Inches(1.0), fill="card", line="line", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, Inches(2.4), Inches(1.4), Inches(0.8), syl, size=30, color="light", font=HEAD, align=PP_ALIGN.CENTER)
text(s, Inches(6.0), Inches(2.35), Inches(0.8), Inches(0.8), "→", size=36, color="gold", font=HEAD)
rect(s, Inches(7.0), Inches(2.0), Inches(5.3), Inches(1.6), fill="hindu", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(7.0), Inches(2.25), Inches(5.3), Inches(1.0), "КОРОВА", size=40, color="white", font=HEAD,
     bold=True, align=PP_ALIGN.CENTER)
text(s, Inches(1.0), Inches(3.55), Inches(5), Inches(0.5), "звуки идут по очереди и исчезают", size=15, color="light", italic=True)
text(s, Inches(7.0), Inches(3.75), Inches(5.3), Inches(0.5), "пратибха — мгновенное схватывание", size=15, color="light", italic=True, align=PP_ALIGN.CENTER)
rect(s, Inches(1.0), Inches(4.55), Inches(11.3), Inches(1.6), fill="card", line="line")
text(s, Inches(1.3), Inches(4.7), Inches(10.8), Inches(1.4),
     [[("Хелен Келлер, 1887: ", {"bold": True, "color": "dark"}), ("вода на ладони и «w-a-t-e-r» на другой — и мир «затрепетал жизнью».", {})]],
     size=21, color="ink")
notes(s, 15, "Спхота, Хелен Келлер, пратибха")

s = base("Три уровня речи", S3, "Бхартрихари ↔ путь молитвы")
levels = [("вайкхари", "произнесённая речь", "устная молитва"),
          ("мадхьяма", "внутренняя речь", "умная молитва"),
          ("пашьянти", "«видящая» речь: слово и смысл ещё одно", "за пределом слов — изумление (Исаак Сирин)")]
for i, (t, d, p) in enumerate(levels):
    y = Inches(2.05) + i * Inches(1.35)
    shade = ["card", "line", "hindu"][i]
    rect(s, Inches(0.9) + i * Inches(0.35), y, Inches(5.6) - i * Inches(0.7), Inches(1.15), fill=shade,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    col = "white" if i == 2 else "dark"
    text(s, Inches(1.1) + i * Inches(0.35), y + Inches(0.1), Inches(5.2) - i * Inches(0.7), Inches(0.5), t, size=22,
         color=col, font=HEAD, bold=True)
    text(s, Inches(1.1) + i * Inches(0.35), y + Inches(0.58), Inches(5.2) - i * Inches(0.7), Inches(0.5), d, size=15,
         color=col)
    text(s, Inches(6.9), y + Inches(0.3), Inches(0.5), Inches(0.5), "↔", size=24, color="gold", font=HEAD)
    text(s, Inches(7.5), y + Inches(0.3), Inches(5.0), Inches(0.8), p, size=19, color="patristic")
text(s, Inches(0.9), Inches(6.15), Inches(11.6), Inches(0.5),
     "Логос — Лицо, «Слово стало плотью»; шабда-брахман безличен", size=16, color="mid", italic=True)
notes(s, 16, "Вайкхари, мадхьяма, пашьянти; Витгенштейн и отцы")

s = base("Абхинавагупта: узнавание", S3, "Кашмир, около 1000 г.")
portrait(s, "abhinavagupta", Inches(0.9), Inches(2.1), Inches(2.0))
for i, (t, d) in enumerate([("пратьябхиджня", "освобождение — узнать то, чем всегда был"),
                            ("спанда", "реальность — пульсация самосознания"),
                            ("раса", "искусство стирает границу «я» и «это»")]):
    y = Inches(2.05) + i * Inches(1.3)
    text(s, Inches(3.5), y, Inches(3.4), Inches(0.6), t, size=26, color="hindu", font=HEAD, bold=True)
    text(s, Inches(6.9), y + Inches(0.07), Inches(5.7), Inches(0.9), d, size=20, color="ink")
text(s, Inches(3.5), Inches(5.95), Inches(9), Inches(0.5),
     "А вы помните мгновение, когда вас «не было», а было только это?", size=17, color="mid", italic=True)
notes(s, 17, "Абхинавагупта: пратьябхиджня, спанда, раса")

# ── IV. Madhyamaka ──────────────────────────────────────────
divider(18, "IV", "Мадхьямака", "Пустота и взаимозависимое возникновение",
        ["nagarjuna", "chandrakirti", "nagasena"], S4)

s = base(None, S4)
quote(s, "Взаимозависимое возникновение мы называем пустотой. Это — условное обозначение, и это — срединный путь",
      "Нагарджуна, «Муламадхьямака-карика» 24.18", y=Inches(1.3), size=30)
text(s, Inches(1.25), Inches(5.35), Inches(10.8), Inches(0.8),
     "Огонь — только при топливе; топливо — только при огне. Ничто не существует «от себя».",
     size=19, color="buddhist", italic=True)
notes(s, 19, "Шуньята, пратитья-самутпада, огонь и топливо")

s = base("Две истины", S4, "Нагарджуна")
for i, (h, t, ex, c) in enumerate([
        ("Условная", "самвритти-сатья", "«Передай соль». «Вода кипит при 100°». Слова работают.", "line"),
        ("Высшая", "парамартха-сатья", "Всё пусто от самобытия. «Вещей самих по себе» нет.", "buddhist")]):
    x = Inches(0.7) + i * Inches(6.05)
    rect(s, x, Inches(2.0), Inches(5.8), Inches(3.2), fill="card", line=c if c != "line" else "line", lw=1.5)
    text(s, x + Inches(0.3), Inches(2.2), Inches(5.2), Inches(0.6), h, size=30, color="buddhist" if i else "dark", font=HEAD, bold=True)
    text(s, x + Inches(0.3), Inches(2.85), Inches(5.2), Inches(0.5), t, size=17, color="light", italic=True)
    text(s, x + Inches(0.3), Inches(3.5), Inches(5.2), Inches(1.6), ex, size=21, color="ink")
text(s, Inches(0.7), Inches(5.5), Inches(11.9), Inches(1.0),
     ["Без опоры на условное нельзя указать на высшее.",
      [("Неверно понятая пустота губит, как змея, схваченная не за то место (24.11)", {"italic": True, "color": "mid", "size": 17})]],
     size=20, color="dark")
notes(s, 20, "Две истины, Чандракирти, змея, молчание Будды")

s = base("Колесница и светильник", S4, "«Вопросы царя Милинды»")
parts_ = ["дышло", "колёса", "кузов", "ось", "все вместе?"]
for i, p in enumerate(parts_):
    x = Inches(0.8) + i * Inches(1.25)
    rect(s, x, Inches(2.2), Inches(1.1), Inches(0.9), fill="card", line="line", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, Inches(2.4), Inches(1.1), Inches(0.6), p, size=14, color="dark", align=PP_ALIGN.CENTER)
    text(s, x, Inches(3.15), Inches(1.1), Inches(0.5), "нет", size=14, color="patristic", italic=True, align=PP_ALIGN.CENTER)
text(s, Inches(0.8), Inches(3.8), Inches(6.2), Inches(1.6),
     [[("«Колесница» ", {"bold": True, "color": "buddhist"}), ("— имя частей, собранных так.", {})],
      [("«Нагасена» ", {"bold": True, "color": "buddhist"}), ("— имя пяти скандх.", {})]], size=21, color="ink")
rect(s, Inches(7.5), Inches(2.0), Inches(5.1), Inches(3.9), fill="card", line="line")
text(s, Inches(7.8), Inches(2.15), Inches(4.6), Inches(0.5), "Пламя ночью", size=24, color="dark", font=HEAD, bold=True)
text(s, Inches(7.8), Inches(2.8), Inches(4.6), Inches(3.0),
     ["В первую стражу и в последнюю — то же пламя?",
      [("Ни то же, ни другое.", {"bold": True, "color": "buddhist", "size": 24})],
      [("Парфит: важна непрерывность, а не тождество", {"italic": True, "color": "mid", "size": 16})]], size=19)
notes(s, 21, "Нагасена: колесница и светильник; Парфит")

s = base(None, S4)
portrait(s, "buddhaghosa", Inches(0.9), Inches(1.6), Inches(1.7))
text(s, Inches(3.1), Inches(1.25), Inches(9.3), Inches(0.5), "БУДДХАГХОША · «ВИСУДДХИМАГГА» XVI", size=14, color="gold", bold=True)
text(s, Inches(3.1), Inches(1.85), Inches(9.3), Inches(4.2),
     ["Есть страдание, но нет страдающего.", "Есть действие, но нет деятеля.",
      "Есть угасание, но нет угасшего.", "Есть путь, но нет идущего."],
     size=32, color="dark", font=HEAD, italic=True, spacing=1.25)
text(s, Inches(3.1), Inches(5.9), Inches(9.3), Inches(0.5), "Грамматика без подлежащего — процесс без хозяина",
     size=17, color="buddhist")
notes(s, 22, "Буддхагхоша: страдание без страдающего")

s = base("Что такое «я»?", S4, "Три ответа")
table(s, Inches(0.7), Inches(2.0), Inches(11.9), [
    ["", "Юм", "Нагасена · Буддхагхоша", "Отцы Церкви"],
    ["«Я» —", "пучок восприятий", "условное имя потока скандх", "ипостась, личность"],
    ["Тон вывода", "скептическое беспокойство", "освобождение от цепляния", "общение в любви"],
    ["Отрицают", "субстанцию «я»", "самобытие «я»", "замкнутое, изолированное «я»"],
    ["Итог", "«никто не найден»", "«никого нет»", "«есть Кто-то — в отношении»"],
], [Inches(2.2), Inches(3.2), Inches(3.3), Inches(3.2)],
    header_colors=[None, "western", "buddhist", "patristic"], size=17, head_size=18, rowh=Inches(0.75))
notes(s, 23, "Юм, буддизм, ипостась; гипотеза Гопник")

# ── V. Yogacara ─────────────────────────────────────────────
divider(24, "V", "Йогачара", "Сознание как хранилище: как устроен опыт без «вещей самих по себе»",
        ["asanga", "vasubandhu", "dignaga", "dharmakirti"], S5)

s = base("«Только представление»", S5, "Васубандху, «Двадцать строф»")
text(s, Inches(0.7), Inches(2.0), Inches(5.6), Inches(2.6),
     [[("Беркли: ", {"bold": True, "color": "western"}), ("быть — значит быть воспринимаемым; ум и Бог остаются.", {})],
      "",
      [("Васубандху: ", {"bold": True, "color": "buddhist"}), ("субъект и объект возникают вместе — и оба пусты.", {})],
      "", [("Довод сна: опыт полон — а внешних вещей нет.", {"italic": True, "color": "mid"})]], size=20, color="ink")
rect(s, Inches(6.8), Inches(1.95), Inches(5.8), Inches(4.4), fill="card", line="line")
text(s, Inches(7.1), Inches(2.1), Inches(5.3), Inches(0.5), "Одна река — четыре мира", size=22, color="dark", font=HEAD, bold=True)
for i, (who, what) in enumerate([("люди", "вода"), ("рыбы", "дом и воздух"), ("голодные духи", "гной"), ("небожители", "нектар")]):
    y = Inches(2.8) + i * Inches(0.62)
    text(s, Inches(7.1), y, Inches(2.6), Inches(0.5), who, size=19, color="light")
    text(s, Inches(9.7), y, Inches(2.8), Inches(0.5), what, size=19, color="ink", bold=True)
text(s, Inches(7.1), Inches(5.4), Inches(5.3), Inches(0.8), "«Для чистых всё чисто» (Тит 1:15)", size=17, color="patristic", italic=True)
notes(s, 25, "Виджняптиматра, Беркли, сон, река")

s = base("Восемь видов сознания", S5)
layers = [("1–5", "пять чувств: зрение, слух, обоняние, вкус, осязание", "card"),
          ("6", "мано-виджняна — ум, который собирает и называет", "card"),
          ("7", "манас — постоянное «я», «моё»", "line"),
          ("8", "алая-виджняна — хранилище «семян» (биджа)", "buddhist")]
for i, (n, d, f) in enumerate(layers):
    y = Inches(1.95) + i * Inches(1.05)
    rect(s, Inches(0.9), y, Inches(7.4), Inches(0.9), fill=f, line="line", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    col = "white" if f == "buddhist" else "dark"
    text(s, Inches(1.1), y + Inches(0.17), Inches(0.9), Inches(0.6), n, size=22, color=col, font=HEAD, bold=True)
    text(s, Inches(2.0), y + Inches(0.2), Inches(6.2), Inches(0.6), d, size=18, color=col)
text(s, Inches(8.8), Inches(2.0), Inches(3.8), Inches(4.0),
     [[("Семя ", {"bold": True, "color": "buddhist"}), ("→ опыт", {})], [("Опыт ", {"bold": True, "color": "buddhist"}), ("→ новое семя", {})],
      "", [("«Окуривание»: как ткань впитывает запах благовоний", {"italic": True, "color": "mid", "size": 17})]],
     size=22, color="ink")
notes(s, 26, "Восемь сознаний; алая и биджа")

s = base("Восприятие обусловлено историей", S5, "Самая сильная параллель вечера")
table(s, Inches(0.7), Inches(2.0), Inches(11.9), [
    ["", "Фристон", "Йогачара", "Аскетика"],
    ["Хранилище", "генеративная модель", "алая-виджняна", "навыки души"],
    ["Порождает", "предсказания", "семена прорастают", "помыслы → страсти"],
    ["Обновление", "ошибка предсказания", "«окуривание»", "принятие помысла"],
    ["Цель", "настроить модель", "переворот основания: зеркало", "исцеление: страсть → любовь"],
], [Inches(2.3), Inches(3.2), Inches(3.2), Inches(3.2)],
    header_colors=[None, "western", "buddhist", "patristic"], size=18, head_size=18, rowh=Inches(0.75))
text(s, Inches(0.7), Inches(6.0), Inches(11.9), Inches(0.5), "Одна механика — три разные цели", size=19, color="mid", italic=True)
notes(s, 27, "Фристон ↔ алая ↔ страсти")

s = base("Три природы", S5, "Трисвабхава · Асанга, Васубандху")
for i, (t, sk, ex) in enumerate([("вымышленная", "парикальпита", "змея"), ("зависимая", "паратантра", "верёвка — сплетение волокон"),
                                 ("совершенная", "паринишпанна", "то же — без наложения")]):
    x = Inches(0.7) + i * Inches(4.05)
    rect(s, x, Inches(2.0), Inches(3.8), Inches(2.6), fill="card", line="line")
    text(s, x + Inches(0.25), Inches(2.15), Inches(3.4), Inches(0.6), t, size=25, color="buddhist", font=HEAD, bold=True)
    text(s, x + Inches(0.25), Inches(2.75), Inches(3.4), Inches(0.5), sk, size=16, color="light", italic=True)
    text(s, x + Inches(0.25), Inches(3.35), Inches(3.4), Inches(1.2), ex, size=20, color="ink")
text(s, Inches(0.7), Inches(4.95), Inches(11.9), Inches(1.4),
     [[("Шанкара ", {"bold": True, "color": "hindu"}), ("останавливается на верёвке: основа — Брахман.", {})],
      [("Йогачара ", {"bold": True, "color": "buddhist"}), ("распускает и верёвку: основы нет, есть зависимость.", {})]],
     size=20, color="ink")
notes(s, 28, "Трисвабхава; змея и верёвка ещё раз")

s = base("Восприятие до слова", S5, "Дигнага и Дхармакирти")
card(s, Inches(0.7), Inches(1.95), Inches(3.85), Inches(2.3), "пратьякша", "прямое восприятие — свободно от понятий: единичное «это»", "buddhist")
card(s, Inches(4.75), Inches(1.95), Inches(3.85), Inches(2.3), "викальпа", "понятие «стакан» — общий ярлык, наложенный умом", "buddhist")
card(s, Inches(8.8), Inches(1.95), Inches(3.8), Inches(2.3), "апоха", "слово значит через исключение: «корова» = «не не-корова»", "buddhist")
text(s, Inches(0.7), Inches(4.6), Inches(11.9), Inches(1.8),
     [[("Барретт: ", {"bold": True, "color": "western"}), ("понятия создают эмоции — и это богатство.", {})],
      [("Дхармакирти: ", {"bold": True, "color": "buddhist"}), ("понятия накладываются на ощущение — и это источник ошибки.", {})],
      [("Соссюр: «в языке нет ничего, кроме различий»", {"italic": True, "color": "mid", "size": 17})]], size=20, color="ink")
notes(s, 29, "Пратьякша, викальпа, апоха, свасамведана, Барретт")

# ── VI. Comparison ──────────────────────────────────────────
divider(30, "VI", "Четыре традиции", "Сравнение — и то, что в таблицу не помещается", None, S6)

s = base("Четыре ответа на пять вопросов", S6)
table(s, Inches(0.55), Inches(1.85), Inches(12.25), [
    ["", "Западная философия", "Святоотеческая", "Индуизм", "Буддизм"],
    ["Что есть сознание?", "интенциональный поток; «трудная проблема» квалиа", "нус — «око души», центр в сердце", "Атман = Брахман; чистый свидетель", "поток мгновений; алая — хранилище"],
    ["Роль тела", "от машины к воплощённому разуму", "соучастник молитвы, преображается", "инструмент: асана, пранаяма", "скандха; опора внимательности"],
    ["Язык и опыт", "язык конструирует опыт", "Логос и молитва меняют восприятие", "шабда-брахман: слово — основа", "слова условны и пусты"],
    ["Путь", "рефлексия, редукция, наука", "исихия, умная молитва, обожение", "йога, самадхи, мокша", "медитация, прозрение, нирвана"],
    ["Невыразимо", "квалиа — «каково это»", "Божественная сущность", "турья; «нети, нети»", "шуньята, таковость"],
], [Inches(2.0), Inches(2.6), Inches(2.6), Inches(2.5), Inches(2.55)],
    header_colors=[None, "western", "patristic", "hindu", "buddhist"], size=14, head_size=15, rowh=Inches(0.8))
notes(s, 31, "Большая сравнительная таблица; пример с порезанным пальцем")

s = base("Три соблазнительных отождествления", S6, "Не делайте этого")
for i, (a, b, why) in enumerate([
        ("пустота", "божественный мрак", "одна грамматика отрицаний — противоположная онтология"),
        ("Атман", "нус", "нетварный и безличный — сотворённый и личный"),
        ("мокша", "обожение", "узнавание, что ты всегда был Брахманом, — бесконечный рост в Боге")]):
    y = Inches(1.95) + i * Inches(1.25)
    text(s, Inches(0.7), y, Inches(2.6), Inches(0.6), a, size=26, color="buddhist" if i == 0 else "hindu", font=HEAD, bold=True, align=PP_ALIGN.RIGHT)
    text(s, Inches(3.35), y - Inches(0.02), Inches(0.7), Inches(0.6), "≠", size=30, color="patristic", font=HEAD, bold=True, align=PP_ALIGN.CENTER)
    text(s, Inches(4.05), y, Inches(3.9), Inches(0.6), b, size=26, color="patristic", font=HEAD, bold=True)
    text(s, Inches(8.0), y + Inches(0.05), Inches(4.6), Inches(1.0), why, size=16, color="ink")
rect(s, Inches(0.7), Inches(5.75), Inches(11.9), Inches(0.75), fill="dark")
text(s, Inches(0.7), Inches(5.82), Inches(11.9), Inches(0.6), "Сравнивать не слова, а работу, которую слова делают", size=22,
     color="white", font=HEAD, italic=True, align=PP_ALIGN.CENTER)
notes(s, 32, "Ловушки отождествления; правило; Мертон и Судзуки")

s = base("54 мыслителя на одной ленте", S6, "Лекции I–III")
YEAR = {"western": [1637, 1700, 1670, 1690, 1715, 1745, 1785, 1874, 1913, 1945, 1915, 1956, 1935, 1986, 1975, 1996, 2015, 2010],
        "patristic": [150, 230, 390, 360, 425, 430, 455, 560, 620, 680, 810, 1010, 1050, 1280, 1790, 1855, 1870, 1920],
        "hindu": [200, 470, 800, 1000, 1100, 1925, 1960],
        "buddhist": [-150, 180, 370, 400, 430, 500, 620, 640, 850, 1000, 1240]}
DOM, FR = [-300, 1500, 1800, 2030], [0, 0.56, 0.74, 1]
X0, XW = Inches(2.9), Inches(9.6)


def xof(y):
    for i in range(3):
        if y <= DOM[i + 1]:
            return X0 + int(XW * (FR[i] + (y - DOM[i]) / (DOM[i + 1] - DOM[i]) * (FR[i + 1] - FR[i])))
    return X0 + XW


for a, b, lab, sh in [(-300, 500, "Античность", "card"), (500, 1500, "Средние века", "bg"),
                      (1500, 1900, "Новое время", "card"), (1900, 2030, "XX–XXI", "bg")]:
    rect(s, xof(a), Inches(2.1), xof(b) - xof(a), Inches(3.95), fill=sh)
    text(s, xof(a), Inches(1.8), xof(b) - xof(a), Inches(0.35), lab, size=12, color="light", italic=True, align=PP_ALIGN.CENTER)
R = Inches(0.11)
for li, t in enumerate(["western", "patristic", "hindu", "buddhist"]):
    yc = Inches(2.55) + li * Inches(0.95)
    text(s, Inches(0.5), yc - Inches(0.2), Inches(2.35), Inches(0.4), TRAD_RU[t].replace("Святоотеческая аскетика", "Отцы-аскеты"), size=14, color=t, font=HEAD, bold=True, align=PP_ALIGN.RIGHT)
    rows = []
    for yr in sorted(YEAR[t]):
        x = xof(yr); r = 0
        while r < len(rows) and x - rows[r] < 2 * R + Emu(15000):
            r += 1
        if r == len(rows):
            rows.append(x)
        else:
            rows[r] = x
        rect(s, x - R, yc - R + r * (2 * R + Emu(15000)) - Inches(0.15), 2 * R, 2 * R, fill=t, line="white", lw=0.5, shape=MSO_SHAPE.OVAL)
for yr in [-200, 0, 400, 800, 1200, 1600, 1800, 1900, 2000]:
    text(s, xof(yr) - Inches(0.5), Inches(6.08), Inches(1), Inches(0.3), str(yr) if yr >= 0 else "200 до н.э.", size=11, color="mid", align=PP_ALIGN.CENTER)
notes(s, 33, "Хронология 54 мыслителей")

# ── VII. Silence ────────────────────────────────────────────
divider(34, "VII", "Молчание и практика", "Можно ли понять сознание, не изменив его?",
        ["linji_yixuan", "tilopa", "dogen"], S7)

s = base("Четыре практики", S7)
pr = [("исихия", "безмолвие; ум сводится в сердце; Иисусова молитва; Никифор, Симеон", "patristic", "цель — встреча"),
      ("самадхи", "вершина восьми ступеней; исчезает зазор знающего и знаемого", "hindu", "цель — свидетель"),
      ("ниродха-самапатти", "прекращение восприятия и чувства — ещё не нирвана", "buddhist", "цель — прозрение"),
      ("дзадзэн", "сидеть — уже проявление пробуждения (Догэн)", "buddhist", "цель — не цель")]
for i, (h, d, c, g) in enumerate(pr):
    x = Inches(0.7) + (i % 2) * Inches(6.05)
    y = Inches(1.95) + (i // 2) * Inches(2.25)
    rect(s, x, y, Inches(5.8), Inches(2.0), fill="card", line="line")
    rect(s, x, y, Inches(0.1), Inches(2.0), fill=c)
    text(s, x + Inches(0.35), y + Inches(0.15), Inches(5.2), Inches(0.55), h, size=24, color=c, font=HEAD, bold=True)
    text(s, x + Inches(0.35), y + Inches(0.75), Inches(5.2), Inches(0.8), d, size=16, color="ink")
    text(s, x + Inches(0.35), y + Inches(1.5), Inches(5.2), Inches(0.4), g, size=15, color="mid", italic=True)
notes(s, 35, "Исихия, самадхи, ниродха, дзадзэн; Симеон")

s = base("Зеркало и пыль", S7, "«Сутра помоста»")
for i, (who, verse, c) in enumerate([
        ("Шэньсю", "Тело — дерево пробуждения, ум — подставка ясного зеркала. Протирай его усердно, не давай пыли осесть.", "patristic"),
        ("Хуэйнэн", "У пробуждения нет дерева, у зеркала нет подставки. Изначально нет ни одной вещи — куда сесть пыли?", "buddhist")]):
    x = Inches(0.7) + i * Inches(6.05)
    rect(s, x, Inches(1.95), Inches(5.8), Inches(2.85), fill="card", line=c, lw=1.25)
    text(s, x + Inches(0.3), Inches(2.1), Inches(5.2), Inches(0.5), who, size=24, color=c, font=HEAD, bold=True)
    text(s, x + Inches(0.3), Inches(2.7), Inches(5.2), Inches(2.0), verse, size=20, color="dark", font=HEAD, italic=True)
text(s, Inches(0.7), Inches(5.05), Inches(11.9), Inches(1.4),
     [[("Догэн: ", {"bold": True, "color": "buddhist"}), ("изучать себя — значит забыть себя.", {"italic": True})],
      [("Линьцзи: «встретишь Будду — убей Будду». Тилопа: «…не управляй — отдыхай».", {"size": 17, "color": "mid"})]],
     size=21, color="dark")
notes(s, 36, "Шэньсю и Хуэйнэн, Линьцзи, Тилопа, Догэн")

s = base("Что значит «преображённое восприятие»?", S7)
rows = [("Исихасты", "изменение органа: очищенный нус видит то, что было всегда", "patristic"),
        ("Веданта", "узнавание: всё есть Брахман", "hindu"),
        ("Мадхьямака", "тот же мир, увиденный как зависимый и свободный", "buddhist"),
        ("Йогачара", "переворот основания: хранилище становится зеркалом", "buddhist")]
for i, (h, d, c) in enumerate(rows):
    y = Inches(1.95) + i * Inches(0.82)
    text(s, Inches(0.7), y, Inches(2.6), Inches(0.6), h, size=22, color=c, font=HEAD, bold=True)
    text(s, Inches(3.3), y + Inches(0.05), Inches(9.3), Inches(0.7), d, size=20, color="ink")
rect(s, Inches(0.7), Inches(5.35), Inches(11.9), Inches(1.1), fill="card", line="line")
text(s, Inches(0.95), Inches(5.45), Inches(11.4), Inches(1.0),
     [[("Палама против Варлаама, XIV в.: ", {"bold": True, "color": "patristic"}),
       ("есть восприятие, которое не сводится ни к физическому, ни к психическому.", {})]], size=18, color="ink")
notes(s, 37, "Преображённое восприятие; Палама и Варлаам; нейронаука")

# ── VIII. Synthesis ─────────────────────────────────────────
divider(38, "VIII", "Синтез цикла", "Три вечера, четыре традиции, пятьдесят четыре мыслителя", None, S8)

s = base("В чём согласие", S8, "Диагноз")
ag = [("Опыт имеет сторону «изнутри»", "квалиа · сердце · Пуруша · самоосознавание"),
      ("Восприятие окрашено прошлым", "страсти · вритти · семена алаи · априорные модели"),
      ("Тело — не футляр для ума", "Мерло-Понти · исихастское дыхание · асаны · внимание к телу"),
      ("Язык — инструмент и граница", "строит опыт и может его скрыть")]
for i, (h, d) in enumerate(ag):
    y = Inches(1.95) + i * Inches(1.1)
    text(s, Inches(0.7), y, Inches(0.7), Inches(0.7), str(i + 1), size=32, color="goldl", font=HEAD, bold=True)
    text(s, Inches(1.4), y, Inches(11), Inches(0.5), h, size=24, color="dark", font=HEAD, bold=True)
    text(s, Inches(1.4), y + Inches(0.5), Inches(11), Inches(0.5), d, size=17, color="light", italic=True)
notes(s, 39, "Четыре согласия")

s = base("В чём спор", S8, "Цель — и она важнее диагноза")
table(s, Inches(0.7), Inches(1.95), Inches(11.9), [
    ["Вопрос", "Веданта", "Буддизм", "Отцы", "Запад"],
    ["Сознание", "основа всего", "ни основа, ни продукт", "сотворено для Нетварного", "продукт мозга? открыто"],
    ["«Я»", "Атман — абсолют", "условное имя", "личность в отношении", "от субстанции до «модели себя»"],
    ["Преображение", "узнать Брахман", "угасание жажды", "бесконечный рост в Боге", "лучшая настройка?"],
    ["За пределом языка", "Абсолют без лица", "таковость", "Лицо", "?"],
], [Inches(2.3), Inches(2.4), Inches(2.4), Inches(2.4), Inches(2.4)],
    header_colors=[None, "hindu", "buddhist", "patristic", "western"], size=16, head_size=17, rowh=Inches(0.8))
notes(s, 40, "Четыре спора; почему расхождения важнее")

s = base("Что остаётся открытым", S8)
oq = [("Можно ли понять сознание, не изменив его?", "Варела: нейрофеноменология — опыт от первого лица как данные"),
      ("Как говорить о невыразимом?", "апофатика · «нети, нети» · строфы Нагарджуны · коан"),
      ("Какой язык мы выбираем для своей жизни?", "понятия · помыслы · вторая стрела")]
for i, (q, d) in enumerate(oq):
    y = Inches(1.95) + i * Inches(1.45)
    rect(s, Inches(0.7), y, Inches(11.9), Inches(1.25), fill="card", line="line")
    text(s, Inches(1.0), y + Inches(0.12), Inches(11.3), Inches(0.6), q, size=25, color="dark", font=HEAD, bold=True)
    text(s, Inches(1.0), y + Inches(0.7), Inches(11.3), Inches(0.5), d, size=17, color="light", italic=True)
notes(s, 41, "Три открытых вопроса")

s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, fill="dark")
rect(s, Inches(0.35), Inches(0.35), W - Inches(0.7), H - Inches(0.7), line="gold", lw=0.75)
text(s, 0, Inches(0.8), W, Inches(0.6), "❦", size=30, color="goldl", font=HEAD, align=PP_ALIGN.CENTER)
text(s, Inches(1), Inches(1.5), W - Inches(2), Inches(2.2),
     ["Язык строит опыт —", "и указывает за свои пределы."], size=44, color="white", font=HEAD, bold=True,
     align=PP_ALIGN.CENTER, spacing=1.15)
rect(s, (W - Inches(1.6)) // 2, Inches(3.85), Inches(1.6), Emu(22000), fill="goldl")
text(s, Inches(1.8), Inches(4.2), W - Inches(3.6), Inches(1.6),
     "«Мудрость говорит мне, что я — ничто. Любовь говорит, что я — всё. Между ними течёт моя жизнь»",
     size=24, color="pale", font=HEAD, italic=True, align=PP_ALIGN.CENTER)
text(s, 0, Inches(5.85), W, Inches(0.5), "— Нисаргадатта Махарадж, «Я есть То»", size=16, color="goldl", align=PP_ALIGN.CENTER)
notes(s, 42, "Итоговый тезис цикла")

s = base("Вопросы для обсуждения", S8, "Спасибо!")
dq = ["Какая пара понятий из сегодняшней лекции кажется вам действительно близкой — и в чём их различие?",
      "Может ли наука о сознании обойтись без знания «изнутри практики»?",
      "Если язык создаёт опыт, за что отвечает тот, кто говорит?",
      "Какой из четырёх языков ближе вам теперь — и изменился ли ответ за вечер?"]
for i, q in enumerate(dq):
    y = Inches(1.95) + i * Inches(0.95)
    text(s, Inches(0.7), y, Inches(0.6), Inches(0.6), str(i + 1), size=26, color="goldl", font=HEAD, bold=True)
    text(s, Inches(1.35), y + Inches(0.05), Inches(11.2), Inches(0.9), q, size=20, color="dark")
text(s, Inches(0.7), Inches(5.95), Inches(11.9), Inches(0.5),
     "kagort.github.io/language-of-experience — словарь, карта, хронология, персоналии", size=16, color="light", italic=True)
notes(s, 43, "Вопросы для обсуждения")

# ── save ────────────────────────────────────────────────────
assert len(prs.slides) == 43, len(prs.slides)
prs.save(HERE / "lecture3-slides.pptx")

SECTION_OF = {}
for sec, rng in [(1, range(1, 6)), (2, range(6, 13)), (3, range(13, 18)), (4, range(18, 24)),
                 (5, range(24, 30)), (6, range(30, 34)), (7, range(34, 38)), (8, range(38, 44))]:
    for n in rng:
        SECTION_OF[n] = sec
FILES = {1: "01-intro", 2: "02-hinduism-vedanta", 3: "03-hinduism-language", 4: "04-buddhism-madhyamaka",
         5: "05-buddhism-yogacara", 6: "06-comparison", 7: "07-silence-practice", 8: "08-conclusion"}
lines = ["# Лекция III — план слайдов", "",
         "Сгенерировано `build_slides.py`. Колонка «Раздел сайта» — файл в `_sections3/`.", "",
         "| № | Слайд | Раздел сайта |", "|---|-------|--------------|"]
for n, t in OUTLINE:
    lines.append(f"| {n} | {t} | `{FILES[SECTION_OF[n]]}.md` |")
(HERE / "lecture3-slides-outline.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
print("ok:", len(prs.slides), "slides")
